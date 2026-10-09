"""Read monthly denominators without inventing individual working dates."""
import calendar
import re
from openpyxl import load_workbook
from app.domain.statistics import months


def parse_sheet(rows, sheet, start, end):
    values = {}
    header = False
    for row_number, row in enumerate(rows, 1):
        month, count = (list(row) + [None, None])[:2]
        if not header:
            if month == 'Month' and count == 'Working days':
                header = True
            continue
        if month is None and count is None:
            continue
        if not isinstance(month, str) or not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', month):
            raise ValueError(f'{sheet}!A{row_number}: invalid month {month!r}')
        if month in values:
            raise ValueError(f'{sheet}!A{row_number}: duplicate month {month}')
        if count is not None:
            year, number = map(int, month.split('-'))
            if (isinstance(count, bool) or not isinstance(count, (int, float))
                    or not 0 <= count <= calendar.monthrange(year, number)[1]
                    or count != int(count)):
                raise ValueError(f'{sheet}!B{row_number}: invalid working-day count {count!r}')
            count = int(count)
        values[month] = (count, row_number)
    if not header:
        raise ValueError(f'{sheet}: expected Month / Working days column headers')
    return [(month, *values.get(month, (None, None))) for month in months(start, end)]


def read_working_days(path, config):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        result = {}
        for court in config['courts']:
            sheet = config['monthly_working_days']['sheets'][court['id']]
            if sheet not in book.sheetnames:
                raise ValueError(f'{path.name}: missing sheet {sheet}')
            result[court['id']] = parse_sheet(book[sheet].values, sheet, config['start'], config['end'])
        return result
    finally:
        book.close()
