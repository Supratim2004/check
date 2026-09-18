"""Discover and read district workbooks. No court statistics are embedded here."""
from datetime import datetime
from functools import lru_cache
from pathlib import Path
import math
import re
from openpyxl import load_workbook

class DataError(ValueError):
    """A source workbook cannot safely be served."""

HEADERS = {
    "Case type": "case_type", "Side": "side", "Filings": "filings",
    "Mean/mo": "mean", "Median/mo": "median", "Max/mo": "max",
    "Daily rate": "daily_rate", "N": "n", "Mean": "mean",
    "Median": "median", "75th": "p75", "90th": "p90", "95th": "p95",
}
SHEETS = {
    "Arrival": "arrivals", "Gap_workdays": "gap", "Hearings_X": "hearings_per_case",
    "Hearing_rate_monthly": "hearing_rate_monthly", "Disposal_workdays": "disposal",
}
CASEWISE_HEADERS = {
    "Case type": "case_type", "Side": "side", "N (wd)": "n_wd",
    "WD mean": "wd_mean", "WD median": "wd_median", "WD 75th": "wd_p75",
    "WD 90th": "wd_p90", "WD 95th": "wd_p95", "N (mo)": "n_mo",
    "Mo mean": "mo_mean", "Mo median": "mo_median", "Mo 75th": "mo_p75",
    "Mo 90th": "mo_p90", "Mo 95th": "mo_p95",
}

def workbook_files(folder):
    return sorted(p for p in folder.glob("Monthly_*.xlsx") if not p.name.startswith("~$"))

def district_folders(root):
    root = Path(root)
    if not root.is_dir():
        raise DataError(f"Data folder does not exist: {root}")
    return {p.name: p for p in sorted(root.iterdir())
            if p.is_dir() and not p.name.startswith('.') and workbook_files(p)}

def number(value, context, optional=False):
    if value is None or value in ("—", "-", "", "–"):
        if optional:
            return None
        raise DataError(f"{context}: missing numeric value")
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
        raise DataError(f"{context}: expected a finite non-negative number, got {value!r}")
    return value

def table(book, name):
    if name not in book.sheetnames:
        raise DataError(f"missing sheet {name}")
    rows = iter(book[name].values)
    for row in rows:
        if row[0] == "Case type":
            headers = list(row)
            break
    else:
        raise DataError(f"{name}: missing Case type header")
    expected = (["Case type", "Side", "Filings", "Mean/mo", "Median/mo", "Max/mo", "Daily rate"]
                if name == "Arrival" else
                ["Case type", "Side"] + ([] if name == "Hearing_rate_monthly" else ["N"])
                + ["Mean", "Median", "75th", "90th", "95th"])
    if any(h not in headers for h in expected):
        raise DataError(f"{name}: expected columns {expected}")
    result, seen = [], set()
    for row in rows:
        if all(v is None for v in row):
            continue
        item = {}
        for header in expected:
            value = row[headers.index(header)]
            key = HEADERS[header]
            if key in ("case_type", "side"):
                if not isinstance(value, str) or not value.strip():
                    raise DataError(f"{name}: missing {header}")
                item[key] = value.strip()
            else:
                item[key] = number(value, f"{name}/{row[0]}/{header}",
                                   optional=(name != "Arrival" or key == "daily_rate"))
        if item["case_type"] in seen:
            raise DataError(f"{name}: duplicate case type {item['case_type']}")
        seen.add(item["case_type"])
        result.append(item)
    return result


def casewise_table(book):
    """Read the per-case hearing rate distributions in working days and months."""
    name = "Hearing_rate_casewise"
    if name not in book.sheetnames:
        raise DataError(f"missing sheet {name}")
    rows = iter(book[name].values)
    for row in rows:
        if row[0] == "Case type":
            headers = list(row)
            break
    else:
        raise DataError(f"{name}: missing Case type header")
    expected = list(CASEWISE_HEADERS)
    if any(header not in headers for header in expected):
        raise DataError(f"{name}: expected columns {expected}")
    result, seen = [], set()
    for row in rows:
        if all(value is None for value in row):
            continue
        item = {}
        for header in expected:
            value = row[headers.index(header)]
            key = CASEWISE_HEADERS[header]
            if key in ("case_type", "side"):
                if not isinstance(value, str) or not value.strip():
                    raise DataError(f"{name}: missing {header}")
                item[key] = value.strip()
            else:
                item[key] = number(value, f"{name}/{row[0]}/{header}", optional=True)
        if item["case_type"] in seen:
            raise DataError(f"{name}: duplicate case type {item['case_type']}")
        seen.add(item["case_type"])
        result.append(item)
    return result

def read_court(path):
    try:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            definitions = dict((r[0], r[1]) for r in book["Definitions"].values if r[0])
            title = next(iter(definitions))
            court = {"full_name": title.split(" — ")[0], "source_file": path.name}
            court.update({key: table(book, sheet) for sheet, key in SHEETS.items()})
            court["hearing_rate_casewise"] = casewise_table(book)
            if not court["arrivals"]:
                raise DataError("Arrival has no case types")
            court["arrivals"].sort(key=lambda r: (-r["filings"], r["case_type"]))
            cutoff = re.search(r"\d{4}-\d{2}-\d{2}", definitions.get("Court cut-off", ""))
            court["listing_cutoff"] = cutoff.group() if cutoff else None
            extracted = re.search(r"Data extracted ([^;]+)", definitions.get("Scope / window", ""))
            court["extracted_on"] = extracted.group(1).strip() if extracted else None
            panels = [s for s in book.sheetnames if s.startswith("Arrival_panel_")]
            if len(panels) != 1:
                raise DataError("expected one Arrival_panel_* sheet")
            rows = iter(book[panels[0]].values)
            header = next(rows)
            if header[0] != "Month" or len(set(header)) != len(header):
                raise DataError("arrival panel needs Month and unique case-type headers")
            values = {}
            for row in rows:
                if all(v is None for v in row):
                    continue
                month = row[0]
                if not isinstance(month, str) or not re.fullmatch(r"\d{4}-\d{2}", month):
                    raise DataError(f"invalid panel month {month!r}")
                datetime.strptime(month, "%Y-%m")
                if month in values:
                    raise DataError(f"duplicate panel month {month}")
                values[month] = [number(v, f"{panels[0]}/{month}", optional=True) for v in row[1:]]
            if not values:
                raise DataError("arrival panel has no months")
            labels = sorted(values)
            court["months"] = len(labels)
            court["window_start"] = datetime.strptime(labels[0], "%Y-%m").strftime("%b %Y")
            court["window_end"] = datetime.strptime(labels[-1], "%Y-%m").strftime("%b %Y")
            court["month_labels"] = labels
            court["arrival_series"] = {name: [values[m][i] for m in labels] for i, name in enumerate(header[1:])}
            rate_panels = [s for s in book.sheetnames if s.startswith("HearingRate_panel_")]
            if len(rate_panels) != 1:
                raise DataError("expected one HearingRate_panel_* sheet")
            rate_rows = iter(book[rate_panels[0]].values)
            rate_header = next(rate_rows, ())
            if (not rate_header or rate_header[0] != "Month" or
                    len(set(rate_header)) != len(rate_header)):
                raise DataError("hearing-rate panel needs Month and unique case-type headers")
            rate_values = {}
            for row in rate_rows:
                if all(value is None for value in row):
                    continue
                month = row[0]
                if not isinstance(month, str) or not re.fullmatch(r"\d{4}-\d{2}", month):
                    raise DataError(f"{rate_panels[0]}: invalid month {month!r}")
                datetime.strptime(month, "%Y-%m")
                if month in rate_values:
                    raise DataError(f"{rate_panels[0]}: duplicate month {month}")
                rate_values[month] = [number(value, f"{rate_panels[0]}/{month}", optional=True)
                                      for value in row[1:]]
            if sorted(rate_values) != labels:
                raise DataError("hearing-rate panel months do not match arrival panel")
            court["hearing_rate_series"] = {
                name: [rate_values[month][i] for month in labels]
                for i, name in enumerate(rate_header[1:])
            }
            working_days = {}
            working_sheet = None
            for sheet in book:
                sheet_rows = iter(sheet.values)
                first = next(sheet_rows, ())
                if tuple(first[:2]) != ("Month", "Working days"):
                    continue
                if working_sheet is not None:
                    raise DataError("multiple working-day sheets")
                working_sheet = sheet.title
                for row in sheet_rows:
                    if all(value is None for value in row):
                        continue
                    month = row[0]
                    if not isinstance(month, str) or not re.fullmatch(r"\d{4}-\d{2}", month):
                        raise DataError(f"{sheet.title}: invalid month {month!r}")
                    datetime.strptime(month, "%Y-%m")
                    if month in working_days:
                        raise DataError(f"{sheet.title}: duplicate month {month}")
                    days = number(row[1], f"{sheet.title}/{month}", optional=True)
                    if days is not None and (days > 31 or int(days) != days):
                        raise DataError(f"{sheet.title}/{month}: invalid working-day count")
                    working_days[month] = days
            court["working_days_source"] = working_sheet
            court["working_days_monthly"] = [working_days.get(month) for month in labels]
            court["arrival_rate_series"] = {
                name: [count / days if count is not None and days is not None and days > 0 else None
                       for count, days in zip(counts, court["working_days_monthly"])]
                for name, counts in court["arrival_series"].items()
            }
            return court
        finally:
            book.close()
    except Exception as exc:
        raise DataError(f"{path.parent.name}/{path.name}: {exc}") from exc

@lru_cache(maxsize=32)
def _load_district(signature, district):
    courts = {}
    for filename, _, _ in signature:
        path = Path(filename)
        code = path.stem.removeprefix("Monthly_")
        code = {"CJ_Sr": "CJ (Sr.)", "CJ_Jr": "CJ (Jr.)"}.get(code, code)
        if code in courts:
            raise DataError(f"{district}: duplicate court code {code}")
        courts[code] = read_court(path)
    labels = sorted({m for c in courts.values() for m in c["month_labels"]})
    dates = sorted({c["extracted_on"] or "Not provided" for c in courts.values()})
    meta = {"district": district, "courts": list(courts), "months": len(labels),
            "window_start": datetime.strptime(labels[0], "%Y-%m").strftime("%b %Y"),
            "window_end": datetime.strptime(labels[-1], "%Y-%m").strftime("%b %Y"),
            "extracted_on": ", ".join(dates)}
    return {"meta": meta, "courts": courts}

def load_district(root, district):
    folders = district_folders(root)
    if district not in folders:
        raise KeyError(district)
    signature = tuple((str(p.resolve()), p.stat().st_mtime_ns, p.stat().st_size)
                      for p in workbook_files(folders[district]))
    return _load_district(signature, district)
